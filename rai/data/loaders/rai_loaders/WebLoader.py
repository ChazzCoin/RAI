import threading
from typing import List
from F import DICT, LIST, DATE
from F.LOG import Log
from openpyxl.styles.builtins import title
from selenium.webdriver.common.by import By
from rai.base.BaseAgents import RaiBaseAgent
from rai.base.BaseFormats import RaiMetadata
from rai.base.BaseLoaders import register_loader
from rai.data.loaders.rai_loaders.BaseLoad import RaiBaseLoader
from rai.data.utilities.DataUtilities import ensure_string_for_chroma
from selenium.common import WebDriverException, TimeoutException, NoSuchElementException
from rai.data.loaders.rai_loaders.BaseDoc import RaiLoaderDocument
from rai.data.web.BodyExtractor import WebBodyExtractor
from rai.data.web.WebDrive import RaiWebDriver, RaiUrl
from rai.data.web.WebExtractor import WebActionExtractor
from rai.data.web.WebModels import WebPageDetails

Log = Log("RaiWebLoader")

@register_loader(name="web")
class RaiWebLoader(RaiWebDriver, RaiBaseLoader):
    cache:[] = []
    documents: List['RaiLoaderDocument'] = []

    def __init__(self, base_url: str, username=None, password=None):
        super(RaiWebLoader, self).__init__()
        self.setup_login(username, password)
        self.url = RaiUrl(base_url)
        self.current_site_name = self.url.site_name
        self.domain_name = self.url.savable_name
        self.to_visit_urls = {self.url}
        self.visited_urls = set()  # track visited
        self.documents = []        # store RaiDocument objects
        self.data_lock = threading.Lock()
        self.scrape_count = 0

    @classmethod
    def pipeline(cls, url: str, page_limit: int = 1, username=None, password=None):
        extractor = cls(url, username=username, password=password)
        extractor.scrape_limit = page_limit
        extractor.crawl()
        return extractor

    def load(self):
        if self.cache:
            Log.i(f"Returning Cached Loader: [ {getattr(self, 'file_path', None)} ]")
            return self.cache
        return self.crawl()

    def generate_metadata(self, content:str):
        Log.i("Generating Metadata.")
        metadata: RaiMetadata = RaiBaseAgent.pipeline(name='metadata', user_prompt=content)
        meta = metadata.model_dump()
        meta = DICT.add_key_value('url', self.current_url, meta)
        meta = DICT.add_key_value('page_title', self.page_title, meta)
        return meta

    def to_documents(self, page: WebPageDetails):
        """ Web Contents Loader """
        Log.i("Creating Content Documents.")
        doc = RaiLoaderDocument(
            page_content=str(page.content),
            metadata=ensure_string_for_chroma(page.metadata)
        )
        self.cache.append(doc)
        """ Table Contents Loader """
        Log.i("Creating Table Documents.")
        for tableItem in LIST.flatten(page.table_objs):
            table_doc = RaiLoaderDocument(
                page_content=str(tableItem),
                metadata=ensure_string_for_chroma(page.metadata)
            )
            self.cache.append(table_doc)
        """ Events Contents Loader """
        Log.i("Creating Event Documents.")
        for eventItem in LIST.flatten(page.events):
            event_doc = RaiLoaderDocument(
                page_content=str(eventItem),
                metadata=ensure_string_for_chroma(page.metadata)
            )
            self.cache.append(event_doc)
        Log.i(f"Document Count: {len(self.cache)}")

    def _scrape_page(self, url: str):
        try:
            self.open(url)

            Log.i("Extracting Page.")
            urls = self.extract_urls()
            images = self.extract_image_urls()
            image_texts = self.extract_text_from_image_urls(images)
            events = self.extract_events()
            table_objs = self.extract_all_tables()

            actions = WebActionExtractor.pipeline(self.driver.page_source)
            body = WebBodyExtractor.pipeline(self.driver.page_source)

            self.add_urls_to_queue(urls)

            """ Web Metadata """
            metadata = self.generate_metadata(body.combined_text)
            page = WebPageDetails(
                url=url,
                title=self.page_title,
                author="RaiWebDriver",
                date=DATE.get_now_month_day_year_str(),
                content=body.combined_text,
                body = body,
                actions = actions,
                urls=urls,
                tags=DICT.get_any(keys=("tags", "keywords"), dic=metadata, default=[]),
                images=images,
                images_content=image_texts,
                tables=table_objs,
                events=events,
                metadata=metadata
            )
            self.pages.append(page)
            self.to_documents(page)
            Log.s(f"Successfully Scraped Page: {url}")
        except (WebDriverException, TimeoutException, NoSuchElementException) as e:
            Log.e(f"Error Scraping {url}: {e}")


    def crawl(self):
        """
        Main loop: pop URLs from to_visit_urls, scrape them, and add discovered links.
        Respects the scrape_limit to avoid infinite crawling.
        """
        while self.to_visit_urls:
            if not self.scrape_limit > 0: break
            if self.scrape_count >= self.scrape_limit: break
            next_url = self.to_visit_urls.pop()
            # If already visited, skip
            if next_url in self.visited_urls: continue
            Log.w(f"Crawling has Begun...")
            self.current_url = next_url
            self.visited_urls.add(next_url)
            self._scrape_page(next_url)
            self.scrape_count += 1

        Log.s("Crawling completed.")
        return self.cache


if __name__ == '__main__':
    email = "jperson@parkcitysoccer.org"
    password = "Philly23!"
    crawler = RaiWebLoader.pipeline("https://www.cnn.com/weather/live-news/fire-los-angeles-california-palisades-ventura-eaton-01-15-25-hnk/index.html", 100, username=None, password=None)
    loader = crawler.load()
    for item in loader:
        print(item.page_content)